"""
Document Parser Module for reading resumes in PDF, DOCX, and TXT formats.
"""

import os
from typing import Union, BinaryIO
import pdfplumber
import docx


class DocumentParser:
    """Parses text from various document formats (PDF, DOCX, TXT)."""

    @staticmethod
    def extract_text_from_pdf(file_source: Union[str, BinaryIO]) -> str:
        """Extracts text content from a PDF file path or file-like object."""
        text = ""
        try:
            with pdfplumber.open(file_source) as pdf:
                for page in pdf.pages:
                    extracted = page.extract_text()
                    if extracted:
                        text += extracted + "\n"
        except Exception as e:
            raise ValueError(f"Failed to parse PDF file: {str(e)}")
        return text

    @staticmethod
    def extract_text_from_docx(file_source: Union[str, BinaryIO]) -> str:
        """Extracts text content from a DOCX file path or file-like object."""
        text = ""
        try:
            doc = docx.Document(file_source)
            for paragraph in doc.paragraphs:
                if paragraph.text:
                    text += paragraph.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    row_text = " ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        text += row_text + "\n"
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX file: {str(e)}")
        return text

    @staticmethod
    def extract_text_from_txt(file_source: Union[str, BinaryIO]) -> str:
        """Extracts text content from a plain text file path or file-like object."""
        if isinstance(file_source, str):
            with open(file_source, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        else:
            return file_source.read().decode('utf-8', errors='ignore')

    @classmethod
    def parse_file(cls, file_source: Union[str, BinaryIO], filename: str = "") -> str:
        """
        Determines file type and extracts raw text.
        Works with file paths or uploaded Streamlit file objects.
        """
        name = filename if filename else (file_source if isinstance(file_source, str) else getattr(file_source, 'name', ''))
        ext = os.path.splitext(name)[1].lower()

        if ext == '.pdf':
            return cls.extract_text_from_pdf(file_source)
        elif ext in ['.docx', '.doc']:
            return cls.extract_text_from_docx(file_source)
        elif ext in ['.txt', '.text', '.md']:
            return cls.extract_text_from_txt(file_source)
        else:
            # Attempt to read as text if unknown extension
            try:
                return cls.extract_text_from_txt(file_source)
            except Exception:
                raise ValueError(f"Unsupported file format: {ext}")
